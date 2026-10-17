"""Daily scraper for Zagreb VSS public job notices.

The state portal keeps every VSS job in Zagreb from the Ministry of Health
or the Ministry of Environmental Protection. Every other source keeps VSS
sanitary-engineer jobs in Zagreb.

Reads RESEND_API_KEY, RECEIVER_EMAIL, and RESEND_FROM_EMAIL from the
environment. Tracks mailed notices in sent_jobs.json next to this file so a
GitHub Actions job can commit that file back to the repository.
"""

from __future__ import annotations

import json
import logging
import os
import re
import ssl
import sys
import time
import warnings
import zipfile
from dataclasses import dataclass
from datetime import date, timedelta
from html import escape, unescape
from io import BytesIO
from pathlib import Path
from urllib.parse import urljoin
from xml.etree import ElementTree

import requests
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

BASE_DIR = Path(__file__).resolve().parent
SENT_PATH = BASE_DIR / "sent_jobs.json"

CITY_LIST_URL = "https://zagreb.hr/objavljeni-natjecaji-oglasi/6501"
CULTURE_LIST_URL = (
    "https://zagreb.hr/javni-natjecaji-za-zaposljavanje-u-ustanovama-u-ku/202452"
)
KINDERGARTEN_LIST_URL = "https://vrtici.zagreb.hr/natjecaji-za-zaposljavanje/122"
SELEKCIJA_API_URL = (
    "https://selekcija.gov.hr/MPUCSZ_Backend_Prod/rest/"
    "Natjecaji/ObjavljeniNatjecajiList/"
)
SELEKCIJA_PORTAL_URL = "https://selekcija.gov.hr/natjecaji/objavljeni-natjecaji"
HOLDING_LIST_URL = (
    "https://www.zgh.hr/karijere/javni-natjecaji-za-zaposljavanje/5699"
)
HEP_OGLASI_URL = "https://eprijave.hep.hr/v1/api/oglasi/"
HEP_NATJECAJI_URL = "https://eprijave.hep.hr/v1/api/natjecaji/"
HEP_PORTAL_URL = "https://eprijave.hep.hr"
CESTE_LIST_URL = "https://hrvatske-ceste.hr/hr/stranice/o-nama/zaposljavanja"
VODE_LIST_URL = "https://www.voda.hr/hr/zaposljavanje"
HZPP_LIST_URL = "https://www.hzpp.hr/hr/tvrtka-2/zaposljavanje"
HZINFRA_LIST_URL = "https://www.hzinfra.hr/category/natjecaj-za-zaposljavanje/"
HZCARGO_LIST_URL = "https://www.hzcargo.hr/zaposljavanje/"
HAC_LIST_URL = "https://www.hac.hr/o-nama/zaposljavanje"
SUME_LIST_URL = "https://www.hrsume.hr/vijesti/"
RESEND_URL = "https://api.resend.com/emails"

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
)
REQUEST_PAUSE_SECONDS = 0.35
REQUEST_TIMEOUT = (5, 5)
MAX_DOCUMENT_BYTES = 8_000_000

CITY_TRIGGERS = (
    "natječaj za prijam u službu",
    "oglas za prijam u službu",
    "natječaj za prijam na rad",
    "oglas za prijam na rad",
    "rješenje o djelomičnoj obustavi",
)
CULTURE_TRIGGERS = ("natječaj", "javni natječaj", "oglas za zapošljavanje")

VSS_RE = re.compile(
    r"visok\w*\s+stru[cč]n\w*\s+sprem\w*"
    r"|\bVSS\b"
    r"|\bmagistar\w*"
    r"|\bmagistric\w*"
    r"|diplomsk\w*(?:\s+\S+){0,4}\s+studij\w*",
    re.IGNORECASE,
)
SANITARY_RE = re.compile(r"sanitarn", re.IGNORECASE)
ENGINEER_RE = re.compile(r"in[zž]enjer", re.IGNORECASE)
ZAGREB_CITY_RE = re.compile(r"\bzagreb(?:a|u|om|e)?\b", re.IGNORECASE)
JOB_HEADING_RE = re.compile(r"(?m)^[ \t]*(?=\d{1,2}\.\s+\S)")
MONTHS = {
    "siječnja": 1,
    "veljače": 2,
    "ožujka": 3,
    "travnja": 4,
    "svibnja": 5,
    "lipnja": 6,
    "srpnja": 7,
    "kolovoza": 8,
    "rujna": 9,
    "listopada": 10,
    "studenoga": 11,
    "prosinca": 12,
}

logger = logging.getLogger("scrape_jobs")


@dataclass(frozen=True)
class Job:
    key: str
    office: str
    title: str
    url: str
    deadline: str
    deadline_date: date | None = None
    code: str = ""


def configure_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )
    logging.getLogger("pypdf").setLevel(logging.ERROR)
    warnings.filterwarnings("ignore", module="pypdf")


def require_env() -> dict[str, str]:
    names = ("RESEND_API_KEY", "RECEIVER_EMAIL", "RESEND_FROM_EMAIL")
    values = {name: os.environ.get(name, "").strip() for name in names}
    missing = [name for name, value in values.items() if not value]
    if missing:
        logger.error("Missing environment variables: %s", ", ".join(missing))
        raise SystemExit(1)
    return values


def make_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": USER_AGENT,
            "Accept-Language": "hr-HR,hr;q=0.9,en;q=0.8",
            "Accept": "text/html,application/xhtml+xml,application/json;q=0.9,*/*;q=0.8",
        }
    )
    retry = Retry(
        total=0,
        connect=0,
        read=0,
        status=0,
        redirect=2,
        allowed_methods=frozenset({"GET"}),
    )
    adapter = HTTPAdapter(max_retries=retry)
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    return session


class SystemTrustAdapter(HTTPAdapter):
    """Use the operating-system certificate store when a site's chain is incomplete."""

    def init_poolmanager(self, connections: int, maxsize: int, block: bool = False, **kwargs: object):
        kwargs["ssl_context"] = ssl.create_default_context()
        return super().init_poolmanager(connections, maxsize, block=block, **kwargs)


def session_with_system_trust(session: requests.Session) -> requests.Session:
    trusted = requests.Session()
    trusted.headers.update(session.headers)
    retry = Retry(
        total=0,
        connect=0,
        read=0,
        status=0,
        redirect=2,
        allowed_methods=frozenset({"GET"}),
    )
    trusted.mount("https://", SystemTrustAdapter(max_retries=retry))
    trusted.mount("http://", HTTPAdapter(max_retries=retry))
    return trusted


BLOCKED_URL_RE = re.compile(r"content\.outlook|c:[\\/]users", re.IGNORECASE)


def is_admin_path(url: str) -> bool:
    return BLOCKED_URL_RE.search(url) is not None


def resolve_url(base: str, href: str) -> str | None:
    raw = href.strip()
    lowered = raw.lower()
    if not raw or lowered.startswith(("javascript:", "mailto:", "file:")):
        return None
    if is_admin_path(raw) or re.match(r"^[a-zA-Z]:[\\/]", raw):
        return None
    joined = urljoin(base, raw)
    if is_admin_path(joined) or re.search(r"(?<![A-Za-z])[A-Za-z]:[\\/]", joined):
        return None
    return joined


def fetch(session: requests.Session, url: str) -> requests.Response:
    if is_admin_path(url):
        logger.warning("Skipping administrative link: %s", url[:160])
        raise requests.RequestException("administrative link")
    try:
        response = session.get(url, timeout=REQUEST_TIMEOUT)
        response.raise_for_status()
    except requests.RequestException as exc:
        logger.warning("Skipping slow or failed link: %s (%s)", url[:160], type(exc).__name__)
        raise
    if "html" in response.headers.get("Content-Type", "") or not response.encoding:
        response.encoding = response.apparent_encoding or "utf-8"
    time.sleep(REQUEST_PAUSE_SECONDS)
    return response


def soup_of(session: requests.Session, url: str) -> BeautifulSoup:
    response = fetch(session, url)
    return BeautifulSoup(response.text, "html.parser")


def html_to_text(fragment: str) -> str:
    text = re.sub(r"(?i)<br\s*/?>", "\n", fragment)
    text = re.sub(r"(?i)</(p|div|li|h\d|tr|strong|td)>", "\n", text)
    text = re.sub(r"<[^>]+>", " ", text)
    text = unescape(text).replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*", "\n", text)
    return text.strip()


def element_text(element: BeautifulSoup) -> str:
    return html_to_text(str(element))


def collapse(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def normalize_key_part(value: str) -> str:
    return collapse(value).lower()


def is_vss(text: str) -> bool:
    return VSS_RE.search(text) is not None


def is_sanitary_engineer(text: str) -> bool:
    """Stem match for sanitarni/sanitarna/sanitarnog and inženjer/inzenjer forms."""
    return SANITARY_RE.search(text) is not None and ENGINEER_RE.search(text) is not None


def matches_vss_sanitary(text: str) -> bool:
    """VSS (or an equivalent higher degree) and a sanitary-engineer title."""
    return is_vss(text) and is_sanitary_engineer(text)


def has_city_trigger(text: str) -> bool:
    """Notices often print NATJEČAJ as spaced capital letters."""
    compact = re.sub(r"(?i)(?<=\b\w) (?=\w\b)", "", text)
    compact = re.sub(r"\s+", " ", compact).lower()
    return any(trigger in compact for trigger in CITY_TRIGGERS)


def parse_dmy(value: str) -> date | None:
    match = re.search(r"(\d{1,2})\.\s*(\d{1,2})\.\s*(\d{4})", value)
    if not match:
        return None
    day, month, year = (int(part) for part in match.groups())
    try:
        return date(year, month, day)
    except ValueError:
        return None


def parse_iso(value: str) -> date | None:
    match = re.search(r"(\d{4})-(\d{2})-(\d{2})", value)
    if not match:
        return None
    year, month, day = (int(part) for part in match.groups())
    try:
        return date(year, month, day)
    except ValueError:
        return None


def parse_croatian_date(value: str) -> date | None:
    match = re.search(
        r"(\d{1,2})\.\s*(" + "|".join(MONTHS) + r")\s*(\d{4})",
        value,
        re.IGNORECASE,
    )
    if not match:
        return None
    day = int(match.group(1))
    month = MONTHS[match.group(2).lower()]
    year = int(match.group(3))
    try:
        return date(year, month, day)
    except ValueError:
        return None


def format_dmy(value: date) -> str:
    return value.strftime("%d.%m.%Y.")


def describe_deadline(text: str, published: date | None) -> tuple[str, date | None]:
    """Return a display string and the application deadline date when known."""
    absolute = re.search(
        r"(?:rok za podnošenje prijava|prijave do|vrijedi do|otvoren do)\s*(?:je)?\s*[:\-]?\s*"
        r"(\d{1,2}\.\s*\d{1,2}\.\s*\d{4}\.?)",
        text,
        re.IGNORECASE,
    )
    if absolute:
        found = parse_dmy(absolute.group(1))
        label = collapse(absolute.group(0))
        return (label, found)

    relative = re.search(
        r"u roku(?:\s+od)?\s+(\d+)\s*(?:\([^)]{0,40}\))?\s*dana",
        text,
        re.IGNORECASE,
    )
    if relative and published is not None:
        days = int(relative.group(1))
        due = published + timedelta(days=days)
        label = (
            f"{format_dmy(due)} ({days} dana od objave {format_dmy(published)})"
        )
        return (label, due)
    if relative:
        return (collapse(relative.group(0)), None)

    prose = re.search(
        r"Rok za podnošenje prijava[^.\n]{0,160}",
        text,
        re.IGNORECASE,
    )
    if prose:
        sentence = collapse(prose.group(0))
        return (sentence, parse_dmy(sentence) or parse_croatian_date(sentence))

    iso = parse_iso(text)
    if iso and len(text.strip()) <= 10:
        return (text.strip(), iso)
    return ("nije naveden", None)


def clean_title(title: str) -> str:
    title = collapse(title)
    title = re.sub(r"^\d{1,2}\.\s*", "", title)
    title = re.split(
        r"\s+[–—-]\s+\d+\s+izvršitelj",
        title,
        maxsplit=1,
        flags=re.IGNORECASE,
    )[0]
    return title.strip(" -–—:.")


def position_title(block: str) -> str:
    before_requirements = re.split(
        r"Potrebno stručno znanje",
        block,
        maxsplit=1,
        flags=re.IGNORECASE,
    )[0]
    lines = [collapse(line) for line in before_requirements.splitlines() if collapse(line)]
    for index, line in enumerate(lines):
        if not re.search(r"\d+\s+izvršitelj", line, re.IGNORECASE):
            continue
        count_only = re.match(r"^[–—\-+\s]*\d+\s+izvršitelj", line, re.IGNORECASE)
        if count_only and index > 0:
            return clean_title(lines[index - 1])
        return clean_title(line)
    flattened = collapse(before_requirements)
    match = re.search(r"\d{1,2}\.\s+(.+)", flattened)
    if match:
        return clean_title(match.group(1))
    return clean_title(flattened)


def pdf_text(payload: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(payload))
        pages: list[str] = []
        for page in reader.pages[:20]:
            try:
                pages.append(page.extract_text() or "")
            except Exception as exc:
                logger.warning("Skipping unreadable PDF page: %s", exc)
                continue
        return "\n".join(pages)
    except Exception as exc:
        logger.warning("Skipping corrupted PDF: %s", exc)
        return ""


def docx_text(payload: bytes) -> str:
    try:
        document = Document(BytesIO(payload))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(part for part in parts if part.strip())
    except Exception as exc:
        logger.warning("Skipping unreadable Word file: %s", exc)
        return ""


def is_downloadable_document(url: str) -> bool:
    lowered = url.lower()
    if not lowered.startswith(("http://", "https://")):
        return False
    if any(host in lowered for host in ("facebook.com", "twitter.com", "linkedin.com")):
        return False
    path = lowered.split("?", 1)[0]
    return path.endswith((".pdf", ".doc", ".docx", ".odt")) or "userdocsimages" in path


def document_text(session: requests.Session, url: str, cache: dict[str, str]) -> str:
    if url in cache:
        return cache[url]
    response = fetch(session, url)
    payload = response.content
    if len(payload) > MAX_DOCUMENT_BYTES:
        raise ValueError(f"document is larger than {MAX_DOCUMENT_BYTES} bytes")
    content_type = response.headers.get("Content-Type", "").lower()
    path = url.lower().split("?", 1)[0]
    try:
        if path.endswith(".pdf") or "pdf" in content_type:
            text = pdf_text(payload)
        elif path.endswith((".docx", ".doc")) or "wordprocessingml" in content_type or "msword" in content_type:
            text = docx_text(payload)
        else:
            parsed = BeautifulSoup(payload, "html.parser")
            for tag in parsed(["script", "style", "noscript"]):
                tag.decompose()
            main = parsed.select_one(
                "article, main, .entry-content, .post-content, .opis, .page-content"
            )
            text = (main or parsed).get_text("\n", strip=True)
    except Exception as exc:
        logger.warning("Skipping unreadable document %s (%s)", url[:160], exc)
        text = ""
    cache[url] = text
    return text


def still_open(deadline_date: date | None, today: date) -> bool:
    return deadline_date is None or deadline_date >= today


def make_job(
    source: str,
    office: str,
    title: str,
    url: str,
    deadline: str,
    deadline_date: date | None,
    today: date,
) -> Job | None:
    office = collapse(office)
    title = collapse(title)
    if not office or not title or not url:
        return None
    if not still_open(deadline_date, today):
        logger.info("Skipping expired posting: %s — %s", office, title)
        return None
    key = f"zagreb:{source}:{url}:{normalize_key_part(title)}"
    return Job(
        key=key,
        office=office,
        title=title.upper(),
        url=url,
        deadline=deadline or "nije naveden",
        deadline_date=deadline_date,
    )


def collect_city(session: requests.Session, today: date) -> list[Job]:
    logger.info("Reading city office list %s", CITY_LIST_URL)
    page = soup_of(session, CITY_LIST_URL)
    jobs: list[Job] = []
    items = page.select("div.listaitem")
    logger.info("City list has %s office notices", len(items))
    for item in items:
        link = item.select_one("h4 a")
        if link is None or not link.get("href"):
            continue
        office = link.get_text(" ", strip=True)
        notice_url = urljoin(CITY_LIST_URL, link["href"])
        if is_admin_path(link["href"]) or is_admin_path(notice_url):
            logger.warning("Skipping administrative link: %s", str(link["href"])[:160])
            continue
        published = parse_dmy(item.get_text(" ", strip=True))
        try:
            notice = soup_of(session, notice_url)
        except requests.RequestException:
            continue
        body = notice.select_one("div.opis")
        if body is None:
            logger.warning("No notice body at %s", notice_url)
            continue
        text = element_text(body)
        if not has_city_trigger(text):
            logger.info("Skipping office page without a competition trigger: %s", notice_url)
            continue
        if published is None:
            published = parse_dmy(text) or parse_croatian_date(text)
        deadline, deadline_date = describe_deadline(text, published)
        blocks = [block for block in JOB_HEADING_RE.split(text) if block.strip()]
        kept = 0
        for block in blocks:
            if not matches_vss_sanitary(block):
                continue
            title = position_title(block)
            if re.match(r"(?i)^(zakona|narodne|narodnih|članka|clanka)\b", title):
                continue
            if hosted_outside_zagreb(block) or (
                not text_has_zagreb_workplace(block) and hosted_outside_zagreb(text)
            ):
                logger.info("Skipping city post outside Zagreb: %s — %s", office, title)
                continue
            job = make_job(
                "gradska-uprava",
                office,
                title,
                notice_url,
                deadline,
                deadline_date,
                today,
            )
            if job is not None:
                jobs.append(job)
                kept += 1
        logger.info("%s: %s VSS sanitary-engineer positions", office, kept)
    return jobs


def culture_title_and_office(label: str, tail: str) -> tuple[str, str]:
    title = re.sub(
        r"(?i)^(?:javni\s+)?natječaj\s+za\s+(?:radno\s+mjesto\s+|izbor\s+i\s+imenovanje\s+)?",
        "",
        label,
    )
    title = re.sub(r"(?i)\s*\(m\s*/\s*ž\)\s*$", "", title).strip(" -–")
    office_match = re.search(
        r"-\s*(.+?)\s*\(\s*\d{1,2}\.\s*\d{1,2}\.\s*\d{4}\.?\s*\)\s*$",
        tail,
    )
    if office_match:
        office = office_match.group(1)
    else:
        office = re.sub(r"\(\s*\d{1,2}\.\s*\d{1,2}\.\s*\d{4}\.?\s*\)", "", tail)
        office = office.strip(" -–")
    if not office:
        office = "Ustanova Grada Zagreba"
    if not title:
        title = label
    return collapse(title), collapse(office)


def collect_culture(session: requests.Session, today: date) -> list[Job]:
    logger.info("Reading culture list %s", CULTURE_LIST_URL)
    page = soup_of(session, CULTURE_LIST_URL)
    body = page.select_one("div.opis")
    if body is None:
        logger.warning("Culture page has no notice list")
        return []
    jobs: list[Job] = []
    cache: dict[str, str] = {}
    links = body.find_all("a", href=True)
    logger.info("Culture list has %s links", len(links))
    for link in links:
        href = str(link.get("href") or "")
        if is_admin_path(href):
            logger.warning("Skipping administrative link: %s", href[:160])
            continue
        label = collapse(link.get_text(" ", strip=True))
        if not any(trigger in label.lower() for trigger in CULTURE_TRIGGERS):
            continue
        tail_parts: list[str] = []
        for sibling in link.next_siblings:
            if getattr(sibling, "name", None) == "br":
                break
            if getattr(sibling, "name", None) == "a":
                break
            if hasattr(sibling, "get_text"):
                tail_parts.append(sibling.get_text(" ", strip=True))
            else:
                tail_parts.append(str(sibling))
        tail = collapse(unescape(" ".join(tail_parts)))
        title, office = culture_title_and_office(label, tail)
        document_url = resolve_url(CULTURE_LIST_URL, link["href"])
        if document_url is None:
            logger.warning("Skipping unusable culture link: %s", link["href"][:120])
            continue
        published = parse_dmy(tail) or parse_croatian_date(tail)
        try:
            document = document_text(session, document_url, cache)
        except requests.RequestException:
            continue
        except (ValueError, zipfile.BadZipFile, OSError, ElementTree.ParseError) as exc:
            logger.warning("Skipping unreadable culture document %s (%s)", document_url[:160], exc)
            continue
        notice_text = f"{title}\n{document}"
        if not matches_vss_sanitary(notice_text):
            continue
        if hosted_outside_zagreb(notice_text):
            logger.info("Skipping culture post outside Zagreb: %s", title)
            continue
        deadline, deadline_date = describe_deadline(document, published)
        if deadline_date is None and published is not None and (today - published).days > 45:
            logger.info("Skipping old culture notice without a deadline: %s", title)
            continue
        job = make_job(
            "kultura",
            office,
            title,
            document_url,
            deadline,
            deadline_date,
            today,
        )
        if job is not None:
            jobs.append(job)
    logger.info("Culture VSS sanitary-engineer positions: %s", len(jobs))
    return jobs


def collect_kindergartens(session: requests.Session, today: date) -> list[Job]:
    logger.info("Reading kindergarten list %s", KINDERGARTEN_LIST_URL)
    page = soup_of(session, KINDERGARTEN_LIST_URL)
    cards = page.select(".news-list-wrapper a.news-item")
    logger.info("Kindergarten list has %s notices", len(cards))
    jobs: list[Job] = []
    cache: dict[str, str] = {}
    for card in cards:
        href = card.get("href")
        if not href:
            continue
        if is_admin_path(str(href)):
            logger.warning("Skipping administrative link: %s", str(href)[:160])
            continue
        office = ""
        date_node = card.select_one(".date")
        title_node = card.select_one("h4")
        if date_node is not None:
            office = date_node.get_text(" ", strip=True)
        title = title_node.get_text(" ", strip=True) if title_node is not None else ""
        detail_url = urljoin(KINDERGARTEN_LIST_URL, href)
        try:
            detail = soup_of(session, detail_url)
        except requests.RequestException:
            continue
        content = detail.select_one(".page-content")
        if content is None:
            logger.warning("Kindergarten detail has no content: %s", detail_url)
            continue
        detail_text = content.get_text("\n", strip=True)
        documents = []
        for anchor in content.find_all("a", href=True):
            document_url = resolve_url(detail_url, anchor["href"]) or ""
            if is_admin_path(anchor["href"]) or is_admin_path(document_url):
                logger.warning("Skipping administrative link: %s", anchor["href"][:160])
                continue
            if is_downloadable_document(document_url):
                documents.append(document_url)
        combined = detail_text
        for document_url in documents:
            try:
                combined += "\n" + document_text(session, document_url, cache)
            except Exception as exc:
                logger.warning("Skipping unreadable kindergarten file %s (%s)", document_url, exc)
                continue
        notice_text = f"{title}\n{combined}"
        if not matches_vss_sanitary(notice_text):
            continue
        if hosted_outside_zagreb(notice_text):
            logger.info("Skipping kindergarten post outside Zagreb: %s", title)
            continue
        deadline, deadline_date = describe_deadline(detail_text, None)
        direct_url = documents[0] if documents else detail_url
        job = make_job(
            "vrtici",
            office,
            title,
            direct_url,
            deadline,
            deadline_date,
            today,
        )
        if job is not None:
            jobs.append(job)
    logger.info("Kindergarten VSS sanitary-engineer positions: %s", len(jobs))
    return jobs


MINISTRY_EMPLOYER_MARKERS = (
    "ministarstvo zdravstva",
    "zaštite okoliša",
)


def is_target_ministry(employer: str) -> bool:
    """Health, or environmental protection, matched on a lowercased substring.

    The portal field is DrzavnoTijeloNaziv. A short marker still matches when
    that name is truncated.
    """
    folded = employer.lower()
    return any(marker in folded for marker in MINISTRY_EMPLOYER_MARKERS)


def collect_selekcija(session: requests.Session, today: date) -> list[Job]:
    logger.info("Reading state portal %s", SELEKCIJA_API_URL)
    filters = [
        {"Property": "MjestoRada", "Operation": "Contains", "Value": "Zagreb"},
        {"Property": "StrucnaSprema", "Operation": "Contains", "Value": "VSS"},
        {"Property": "StatusNaziv", "Operation": "Equals", "Value": "Prijave u tijeku"},
    ]
    jobs: list[Job] = []
    skip = 0
    page_size = 50
    while True:
        response = session.get(
            SELEKCIJA_API_URL,
            params={
                "filters": json.dumps(filters, ensure_ascii=False),
                "sort": "RokZaPrijavuISO8601 desc",
                "top": str(page_size),
                "skip": str(skip),
            },
            timeout=REQUEST_TIMEOUT,
        )
        response.raise_for_status()
        records = response.json().get("Records") or []
        logger.info("State portal page skip=%s returned %s rows", skip, len(records))
        for record in records:
            office = str(record.get("DrzavnoTijeloNaziv") or "").strip()
            if not is_target_ministry(office):
                continue
            education = str(record.get("StrucnaSprema") or "")
            workplace = str(record.get("MjestoRada") or "")
            if not is_vss(education) or not location_is_zagreb(workplace):
                continue
            title = str(
                record.get("RadnoMjestoNaziv")
                or record.get("NazivRadnogMjesta")
                or ""
            ).strip()
            competition_id = str(record.get("NatjecajID") or record.get("ID") or "")
            code = str(record.get("SifraNatjecaja") or "").strip()
            due_text = str(record.get("RokZaPrijavuISO8601") or "").strip()
            due = parse_iso(due_text)
            if not competition_id or not title:
                continue
            deadline = due_text or "nije naveden"
            if due is not None and due < today:
                logger.info("Skipping expired state competition %s", code or competition_id)
                continue
            jobs.append(
                Job(
                    key=f"selekcija:{competition_id}",
                    office=office or "Državno tijelo",
                    title=title.upper(),
                    url=SELEKCIJA_PORTAL_URL,
                    deadline=deadline,
                    deadline_date=due,
                    code=code,
                )
            )
        if len(records) < page_size:
            break
        skip += page_size
        time.sleep(REQUEST_PAUSE_SECONDS)
    logger.info("State portal ministry VSS positions: %s", len(jobs))
    return jobs


OUTCOME_TITLE_RE = re.compile(
    r"odluk|rezultat|ishod|poni[šs]ten|obavijest\w* o izboru|neprihva",
    re.IGNORECASE,
)
SALE_TITLE_RE = re.compile(r"prodaj", re.IGNORECASE)
WORKPLACE_ZAGREB_RE = re.compile(
    r"mjesto(?:m)? rada[^.\n]{0,160}\bzagreb(?:a|u|om|e)?\b"
    r"|lokacij[ae][^.\n]{0,100}\bzagreb(?:a|u|om|e)?\b"
    r"|\bu zagrebu\b"
    r"|\brj zagreb(?:a|u|om|e)?\b"
    r"|\(zagreb(?:a|u|om|e)?\)",
    re.IGNORECASE,
)
WORKPLACE_LINE_RE = re.compile(
    r"mjesto(?:m)?\s+rada\s*[:\-]?\s*([^\n.]{0,180})",
    re.IGNORECASE,
)
SKIP_ATTACHMENT_RE = re.compile(
    r"pkz|prednost|osobnih podat|gdpr|prijavnic|prijavu na",
    re.IGNORECASE,
)
EMPLOYMENT_TITLE_RE = re.compile(r"javni oglas|javni poziv", re.IGNORECASE)


def title_is_closed(title: str) -> bool:
    return OUTCOME_TITLE_RE.search(title) is not None


def text_has_zagreb_workplace(text: str) -> bool:
    return WORKPLACE_ZAGREB_RE.search(text) is not None


def location_is_zagreb(location: str) -> bool:
    return ZAGREB_CITY_RE.search(location) is not None


def hosted_outside_zagreb(text: str) -> bool:
    """True when a workplace line is present and none of them name Zagreb."""
    lines = [match.group(1) for match in WORKPLACE_LINE_RE.finditer(text)]
    if not lines:
        return False
    return not any(location_is_zagreb(line) for line in lines)


def attachment_urls(page: BeautifulSoup, page_url: str) -> list[str]:
    found: list[str] = []
    for link in page.find_all("a", href=True):
        raw_href = str(link.get("href") or "")
        if is_admin_path(raw_href):
            logger.warning("Skipping administrative link: %s", raw_href[:160])
            continue
        label = link.get_text(" ", strip=True)
        if SKIP_ATTACHMENT_RE.search(label) or title_is_closed(label):
            continue
        href = resolve_url(page_url, str(link.get("href") or ""))
        if not href or not is_downloadable_document(href) or href in found:
            continue
        found.append(href)
        if len(found) >= 2:
            break
    return found


def read_notice(session: requests.Session, url: str, cache: dict[str, str]) -> str:
    if is_downloadable_document(url):
        return document_text(session, url, cache)
    page = soup_of(session, url)
    parts = [page.get_text("\n", strip=True)]
    for doc_url in attachment_urls(page, url):
        try:
            parts.append(document_text(session, doc_url, cache))
        except (requests.RequestException, ValueError):
            logger.warning("Could not read attachment %s", doc_url)
    return "\n".join(parts)


def finish_employer_job(
    session: requests.Session,
    source: str,
    office: str,
    title: str,
    url: str,
    today: date,
    published: date | None = None,
    location: str = "",
    known_deadline: str = "",
    known_date: date | None = None,
    details: str | None = None,
) -> Job | None:
    title = collapse(title)
    if not title or not url or title_is_closed(title):
        return None
    if details is None:
        try:
            details = read_notice(session, url, {})
        except (requests.RequestException, ValueError):
            logger.warning("Could not read %s notice %s", source, url)
            return None
    deadline_text, deadline_date = describe_deadline(details, published)
    if known_date is not None:
        deadline_date = known_date
        deadline_text = known_deadline or format_dmy(known_date)
    notice_text = f"{title}\n{details}"
    if location:
        zagreb = location_is_zagreb(location)
    elif hosted_outside_zagreb(details):
        zagreb = False
    else:
        zagreb = text_has_zagreb_workplace(notice_text)
    if not zagreb:
        logger.info("Skipping %s post outside Zagreb: %s", source, title)
        return None
    if not matches_vss_sanitary(f"{notice_text}\n{location}"):
        logger.info("Skipping %s post without VSS sanitary engineer: %s", source, title)
        return None
    if deadline_date is None and published is not None and (today - published).days > 45:
        logger.info("Skipping old %s notice without a deadline: %s", source, title)
        return None
    return make_job(source, office, title, url, deadline_text, deadline_date, today)


def _hep_records(session: requests.Session, url: str) -> list[dict[str, object]]:
    response = fetch(session, url)
    payload = response.json()
    if not isinstance(payload, list):
        logger.warning("HEP %s did not return a list", url)
        return []
    return [item for item in payload if isinstance(item, dict)]


def collect_hep(session: requests.Session, today: date) -> list[Job]:
    logger.info("Reading HEP boards")
    jobs: list[Job] = []
    seen: set[str] = set()
    for url in (HEP_OGLASI_URL, HEP_NATJECAJI_URL):
        for item in _hep_records(session, url):
            notice = item.get("Oglas")
            if not isinstance(notice, dict):
                notice = item
            place = str(notice.get("MJESTO_RADA") or "")
            if not location_is_zagreb(place):
                continue
            title = str(notice.get("RADNO_MJESTO") or "").strip()
            office = str(item.get("NazivPoduzeca") or "HEP grupa").strip()
            education = str(item.get("StrucnaSprema") or "")
            conditions = notice.get("EP_UVJETI_OGLASA") or []
            condition_text = education
            if isinstance(conditions, list):
                condition_text += "\n" + "\n".join(
                    str(row.get("VRIJEDNOST_UVJETA") or "")
                    for row in conditions
                    if isinstance(row, dict)
                )
            description = str(notice.get("OPIS_RADNOG_MJESTA") or "")
            details = f"{condition_text}\n{description}"
            if not matches_vss_sanitary(f"{title}\n{details}"):
                logger.info("Skipping HEP post without VSS sanitary engineer: %s", title)
                continue
            oglas_id = notice.get("OGLAS_ID") or item.get("OGLAS_ID")
            try:
                job_url = f"{HEP_PORTAL_URL}/prijava/{int(float(str(oglas_id)))}"
            except (TypeError, ValueError):
                job_url = f"{HEP_PORTAL_URL}/oglasi"
            due = parse_iso(str(notice.get("TRAJE_DO") or ""))
            deadline = format_dmy(due) if due else "nije naveden"
            job = make_job("hep", office, title, job_url, deadline, due, today)
            if job is None or job.key in seen:
                continue
            seen.add(job.key)
            jobs.append(job)
    logger.info("HEP VSS sanitary-engineer positions: %s", len(jobs))
    return jobs


def collect_hrvatske_ceste(session: requests.Session, today: date) -> list[Job]:
    logger.info("Reading Hrvatske ceste %s", CESTE_LIST_URL)
    page = soup_of(session, CESTE_LIST_URL)
    jobs: list[Job] = []
    for post in page.select("div.post"):
        link = post.select_one("h1 a[href], .post__title a[href]")
        if link is None:
            continue
        title = collapse(link.get_text(" ", strip=True))
        url = resolve_url(CESTE_LIST_URL, str(link.get("href") or ""))
        published = parse_dmy(
            (post.select_one(".post__date") or post).get_text(" ", strip=True).replace("/", ".")
        )
        job = finish_employer_job(
            session,
            "hrvatske-ceste",
            "Hrvatske ceste",
            title,
            url,
            today,
            published=published,
        )
        if job is not None:
            jobs.append(job)
    logger.info("Hrvatske ceste VSS sanitary-engineer positions: %s", len(jobs))
    return jobs


def collect_hrvatske_vode(session: requests.Session, today: date) -> list[Job]:
    logger.info("Reading Hrvatske vode %s", VODE_LIST_URL)
    page = soup_of(session, VODE_LIST_URL)
    jobs: list[Job] = []
    for row in page.select(".views-row"):
        location_node = row.select_one(".views-field-field-lokacija")
        location = collapse(location_node.get_text(" ", strip=True)) if location_node else ""
        if not location_is_zagreb(location):
            continue
        if any(title_is_closed(link.get_text(" ", strip=True)) for link in row.find_all("a")):
            continue
        title_node = row.select_one(".views-field-title")
        title = collapse(title_node.get_text(" ", strip=True)) if title_node else ""
        page_link = None
        pdf_link = None
        for link in row.find_all("a", href=True):
            href = resolve_url(VODE_LIST_URL, str(link.get("href") or ""))
            if not href:
                continue
            if is_downloadable_document(href):
                label = link.get_text(" ", strip=True)
                if pdf_link is None and not SKIP_ATTACHMENT_RE.search(label):
                    pdf_link = href
            elif page_link is None:
                page_link = href
        url = pdf_link or page_link or ""
        published_node = row.select_one(".views-field-created")
        published = parse_dmy(
            published_node.get_text(" ", strip=True) if published_node else ""
        )
        job = finish_employer_job(
            session,
            "hrvatske-vode",
            "Hrvatske vode",
            title,
            url,
            today,
            published=published,
            location=location,
        )
        if job is not None:
            jobs.append(job)
    logger.info("Hrvatske vode VSS sanitary-engineer positions: %s", len(jobs))
    return jobs


def collect_hzpp(session: requests.Session, today: date) -> list[Job]:
    logger.info("Reading HŽ Putnički prijevoz %s", HZPP_LIST_URL)
    page = soup_of(session, HZPP_LIST_URL)
    archive = page.find("h3", string=lambda value: bool(value and "Arhiva" in value))
    job_list = archive.find_previous("ul") if archive else None
    jobs: list[Job] = []
    if job_list is None:
        logger.info("HŽ Putnički prijevoz VSS sanitary-engineer positions: 0")
        return jobs
    for item in job_list.find_all("li", recursive=False):
        text = collapse(item.get_text(" ", strip=True))
        text = re.sub(r"\(ovdje.*?\)", "", text, flags=re.IGNORECASE)
        if title_is_closed(text):
            continue
        document = None
        for link in item.find_all("a", href=True):
            label = link.get_text(" ", strip=True)
            if SKIP_ATTACHMENT_RE.search(label):
                continue
            document = resolve_url(HZPP_LIST_URL, str(link.get("href") or ""))
            break
        title = re.sub(r"^\d{3,5}-\d{2}\s*", "", text).strip()
        job = finish_employer_job(
            session,
            "hzpp",
            "HŽ Putnički prijevoz",
            title,
            document or "",
            today,
        )
        if job is not None:
            jobs.append(job)
    logger.info("HŽ Putnički prijevoz VSS sanitary-engineer positions: %s", len(jobs))
    return jobs


def collect_hzinfra(session: requests.Session, today: date) -> list[Job]:
    logger.info("Reading HŽ Infrastruktura %s", HZINFRA_LIST_URL)
    try:
        page = soup_of(session, HZINFRA_LIST_URL)
    except requests.exceptions.SSLError:
        logger.warning("HŽ Infrastruktura certificate was not in the default bundle; using the system store")
        session = session_with_system_trust(session)
        try:
            page = soup_of(session, HZINFRA_LIST_URL)
        except requests.exceptions.SSLError:
            logger.warning("Skipping HŽ Infrastruktura: certificate could not be verified")
            return []
    jobs: list[Job] = []
    for article in page.select("article"):
        link = article.select_one("h2 a[href]")
        if link is None:
            continue
        title = collapse(link.get_text(" ", strip=True))
        published = parse_croatian_date(article.get_text(" ", strip=True))
        job = finish_employer_job(
            session,
            "hzinfra",
            "HŽ Infrastruktura",
            title,
            str(link.get("href") or ""),
            today,
            published=published,
        )
        if job is not None:
            jobs.append(job)
    logger.info("HŽ Infrastruktura VSS sanitary-engineer positions: %s", len(jobs))
    return jobs


def collect_hzcargo(session: requests.Session, today: date) -> list[Job]:
    logger.info("Reading HŽ Cargo %s", HZCARGO_LIST_URL)
    page = soup_of(session, HZCARGO_LIST_URL)
    jobs: list[Job] = []
    seen: set[str] = set()
    for link in page.find_all("a", href=True):
        href = str(link.get("href") or "")
        if "/zaposljavanje/" not in href:
            continue
        url = resolve_url(HZCARGO_LIST_URL, href)
        if not url:
            continue
        url = url.rstrip("/")
        if url.rstrip("/") == HZCARGO_LIST_URL.rstrip("/"):
            continue
        if "prednosti-pri-zaposljavanju" in url or url in seen:
            continue
        title = collapse(link.get_text(" ", strip=True))
        if not title or title.lower() in {"zapošljavanje", "hrvatski"}:
            continue
        seen.add(url)
        published = parse_croatian_date(title)
        if published is not None and (today - published).days > 45:
            continue
        title = re.sub(
            r"\s+\d{1,2}\.\s+(?:" + "|".join(MONTHS) + r")\s+\d{4}\.?$",
            "",
            title,
            flags=re.IGNORECASE,
        ).strip(" –-")
        job = finish_employer_job(
            session,
            "hzcargo",
            "HŽ Cargo",
            title,
            url,
            today,
            published=published,
        )
        if job is not None:
            jobs.append(job)
    logger.info("HŽ Cargo VSS sanitary-engineer positions: %s", len(jobs))
    return jobs


def collect_hac(session: requests.Session, today: date) -> list[Job]:
    logger.info("Reading HAC %s", HAC_LIST_URL)
    page = soup_of(session, HAC_LIST_URL)
    jobs: list[Job] = []
    for link in page.select("a.frame-white"):
        title_node = link.find("h2")
        title = collapse(title_node.get_text(" ", strip=True) if title_node else "")
        card = collapse(link.get_text(" ", strip=True))
        deadline_text, deadline_date = describe_deadline(card, None)
        job = finish_employer_job(
            session,
            "hac",
            "Hrvatske autoceste",
            title,
            str(link.get("href") or ""),
            today,
            known_deadline=deadline_text,
            known_date=deadline_date if deadline_text != "nije naveden" else None,
        )
        if job is not None:
            jobs.append(job)
    logger.info("HAC VSS sanitary-engineer positions: %s", len(jobs))
    return jobs


def collect_hrsume(session: requests.Session, today: date) -> list[Job]:
    logger.info("Reading Hrvatske šume %s", SUME_LIST_URL)
    page = soup_of(session, SUME_LIST_URL)
    jobs: list[Job] = []
    seen: set[str] = set()
    for article in page.select("article"):
        link = article.find("a", href=True)
        if link is None:
            continue
        title = collapse(link.get_text(" ", strip=True))
        if not EMPLOYMENT_TITLE_RE.search(title) or title_is_closed(title):
            continue
        if SALE_TITLE_RE.search(title):
            continue
        url = resolve_url(SUME_LIST_URL, str(link.get("href") or ""))
        if not url or url in seen:
            continue
        seen.add(url)
        published = None
        dated = re.search(r"/(\d{4})/(\d{2})/(\d{2})/", url)
        if dated:
            try:
                published = date(int(dated.group(1)), int(dated.group(2)), int(dated.group(3)))
            except ValueError:
                published = None
        job = finish_employer_job(
            session,
            "hrsume",
            "Hrvatske šume",
            title,
            url,
            today,
            published=published,
        )
        if job is not None:
            jobs.append(job)
    logger.info("Hrvatske šume VSS sanitary-engineer positions: %s", len(jobs))
    return jobs


def scrape_holding(sent_ids: set[str]) -> list[Job]:
    """Open VSS sanitary-engineer posts on the Zagrebački Holding board."""
    session = make_session()
    today = date.today()
    logger.info("Reading Holding board %s", HOLDING_LIST_URL)
    page = soup_of(session, HOLDING_LIST_URL)
    jobs: list[Job] = []
    seen: set[str] = set()
    for row in page.select("table.table tr"):
        cells = row.find_all("td")
        if len(cells) < 5:
            continue
        link = cells[1].find("a", href=True)
        if link is None:
            continue
        url = resolve_url(HOLDING_LIST_URL, str(link.get("href") or ""))
        if not url:
            continue
        title = collapse(link.get_text(" ", strip=True))
        if not title or title.lower().startswith("napomena"):
            continue
        full = collapse(cells[1].get_text(" ", strip=True))
        branch = full[len(title):].strip(" -\u00a0–—") if full.lower().startswith(title.lower()) else full
        office = branch or "Zagrebački holding"
        location = collapse(cells[-1].get_text(" ", strip=True))
        if not location_is_zagreb(location):
            continue
        raw_deadline = collapse(cells[3].get_text(" ", strip=True))
        deadline_date = parse_dmy(raw_deadline)
        deadline_text = raw_deadline or "nije naveden"
        key = f"zagreb:holding:{url}:{normalize_key_part(title)}"
        if key in sent_ids or key in seen:
            continue
        try:
            details = document_text(session, url, {})
        except requests.RequestException:
            logger.warning("Could not read Holding notice %s", url)
            continue
        if not matches_vss_sanitary(f"{title}\n{details}"):
            logger.info("Skipping Holding post without VSS sanitary engineer: %s", title)
            continue
        job = make_job(
            "holding",
            office,
            title,
            url,
            deadline_text,
            deadline_date,
            today,
        )
        if job is None:
            continue
        seen.add(job.key)
        jobs.append(job)
    logger.info("Holding VSS sanitary-engineer positions: %s", len(jobs))
    return jobs


def collect_jobs(
    session: requests.Session,
    today: date,
    sent_ids: set[str],
) -> list[Job]:
    collectors = (
        collect_city,
        collect_culture,
        collect_kindergartens,
        collect_selekcija,
        collect_hep,
        collect_hrvatske_ceste,
        collect_hrvatske_vode,
        collect_hzpp,
        collect_hzinfra,
        collect_hzcargo,
        collect_hac,
        collect_hrsume,
    )
    jobs: list[Job] = []
    failures = 0
    for collector in collectors:
        try:
            jobs.extend(collector(session, today))
        except Exception:
            failures += 1
            logger.exception("Source failed: %s", collector.__name__)
    try:
        jobs.extend(scrape_holding(sent_ids))
    except Exception:
        failures += 1
        logger.exception("Source failed: scrape_holding")
    if failures == len(collectors) + 1:
        logger.error("Every source failed")
        raise SystemExit(1)
    unique: dict[str, Job] = {}
    for job in jobs:
        unique.setdefault(job.key, job)
    return list(unique.values())


def load_sent_ids(path: Path) -> set[str]:
    if not path.exists():
        return set()
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        logger.exception("Could not read %s", path)
        raise SystemExit(1) from None
    if isinstance(payload, list):
        return {str(item) for item in payload}
    ids = payload.get("ids", [])
    if not isinstance(ids, list):
        logger.error("%s does not contain an ids list", path)
        raise SystemExit(1)
    return {str(item) for item in ids}


def save_sent_ids(path: Path, ids: set[str]) -> None:
    payload = {"ids": sorted(ids)}
    temporary = path.with_suffix(".json.tmp")
    temporary.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)


def render_messages(jobs: list[Job]) -> tuple[str, str]:
    ordered = sorted(jobs, key=lambda job: (job.office.lower(), job.title.lower()))
    text_lines = [
        "Novi oglasi za radna mjesta VSS u Zagrebu.",
        "Državni portal: sva VSS mjesta Ministarstva zdravstva i Ministarstva zaštite okoliša u Zagrebu. Ostali izvori: sanitarni inženjer.",
        "",
    ]
    html_parts = [
        "<!DOCTYPE html>",
        "<html><body>",
        "<p>Novi oglasi za radna mjesta VSS u Zagrebu.<br>"
        "Državni portal: sva VSS mjesta Ministarstva zdravstva i Ministarstva zaštite okoliša u Zagrebu. Ostali izvori: sanitarni inženjer.</p>",
    ]
    for job in ordered:
        text_lines.append(format_job_text(job))
        text_lines.append(f"  Rok: {job.deadline}")
        text_lines.append("")
        html_parts.append(format_job_html(job))
    html_parts.append("</body></html>")
    return "\n".join(text_lines).strip() + "\n", "\n".join(html_parts)


def format_job_text(job: Job) -> str:
    if job.key.startswith("selekcija:"):
        code = job.code or "nije navedena"
        return (
            f"• {job.office.upper()} - {job.title} "
            f"(Link: {SELEKCIJA_PORTAL_URL} | Šifra: {code} )"
        )
    return f"• {job.office} - {job.title} (Link: {job.url})"


def format_job_html(job: Job) -> str:
    if job.key.startswith("selekcija:"):
        code = job.code or "nije navedena"
        return (
            "<p>• {office} - {title} "
            "(Link: <a href=\"{portal}\">{portal}</a> | "
            "Šifra: <strong>{code}</strong>&nbsp;)"
            "<br>Rok: {deadline}</p>".format(
                office=escape(job.office.upper()),
                title=escape(job.title),
                portal=escape(SELEKCIJA_PORTAL_URL, quote=True),
                code=escape(code),
                deadline=escape(job.deadline),
            )
        )
    return (
        "<p>• {office} - {title} (Link: <a href=\"{url}\">{url}</a>)"
        "<br>Rok: {deadline}</p>".format(
            office=escape(job.office),
            title=escape(job.title),
            url=escape(job.url, quote=True),
            deadline=escape(job.deadline),
        )
    )


def send_email(jobs: list[Job], env: dict[str, str], today: date) -> None:
    if jobs:
        text_body, html_body = render_messages(jobs)
        subject = f"Novi VSS natječaji u Zagrebu ({len(jobs)})"
    else:
        status = (
            "Provjera je uspješno izvršena. Danas nema novih otvorenih VSS "
            "natječaja na praćenim stranicama."
        )
        text_body = status + "\n"
        html_body = (
            "<!DOCTYPE html><html><body><p>"
            + escape(status)
            + "</p></body></html>"
        )
        subject = f"Nema novih VSS Natječaja - Zagreb ({today.strftime('%d.%m.%Y.')})"
    response = requests.post(
        RESEND_URL,
        headers={
            "Authorization": f"Bearer {env['RESEND_API_KEY']}",
            "Content-Type": "application/json",
        },
        json={
            "from": env["RESEND_FROM_EMAIL"],
            "to": [env["RECEIVER_EMAIL"]],
            "subject": subject,
            "html": html_body,
            "text": text_body,
        },
        timeout=REQUEST_TIMEOUT,
    )
    if response.status_code >= 400:
        logger.error("Resend rejected the email: %s %s", response.status_code, response.text)
        raise SystemExit(1)
    message_id = ""
    try:
        message_id = str(response.json().get("id", ""))
    except ValueError:
        logger.error("Resend returned a response that is not JSON")
        raise SystemExit(1) from None
    if jobs:
        logger.info("Sent %s jobs via Resend (%s)", len(jobs), message_id or "no id")
    else:
        logger.info("Sent empty-result status via Resend (%s)", message_id or "no id")


def main() -> None:
    configure_logging()
    env = require_env()
    today = date.today()
    sent_ids = load_sent_ids(SENT_PATH)
    logger.info("Already sent: %s", len(sent_ids))
    jobs = collect_jobs(make_session(), today, sent_ids)
    fresh = [job for job in jobs if job.key not in sent_ids]
    logger.info("Matching jobs: %s, new: %s", len(jobs), len(fresh))
    send_email(fresh, env, today)
    if not fresh:
        logger.info("No new jobs. Status email was sent.")
        return
    save_sent_ids(SENT_PATH, sent_ids | {job.key for job in fresh})
    logger.info("Updated %s", SENT_PATH)


if __name__ == "__main__":
    try:
        main()
    except requests.RequestException:
        logger.exception("HTTP request failed")
        sys.exit(1)
